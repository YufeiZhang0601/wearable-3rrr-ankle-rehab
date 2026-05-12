#!/usr/bin/env python3
"""Append the closing-status section to the Roar.docx speech draft.

Matches the existing bilingual layout (Heading 2 section title, Heading 3
language sub-headers, normal paragraphs underneath). Saves a copy of the
updated document into the repo at ``docs/roar_speech_draft.docx``; if the
``--in-place`` flag is given, the original on disk is also updated.
"""

import argparse
import os
import shutil

import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE = r"d:\Columbia\AMRresearch\Roar.docx"
REPO_COPY = os.path.join(ROOT, "docs", "roar_speech_draft.docx")

ZH_HEADING = u"5. \u9879\u76ee\u72b6\u6001\u4e0e\u5f00\u653e\u5ef6\u7eed"  # 5. 项目状态与开放延续

ZH_BLOCKS = [
    u"\u8fd9\u662f\u6211\u7684\u7ed3\u4e1a\u9879\u76ee\u3002\u867d"
    u"\u7136\u5b9e\u7269\u5c42\u9762\u8fd8\u6ca1\u6709\u5b8c\u5168"
    u"\u8dd1\u901a\uff0c\u4f46\u6211\u4eec\u5df2\u7ecf\u4ea4\u4ed8"
    u"\u4e86\u4e09\u4ef6\u5177\u4f53\u7684\u4e1c\u897f\uff1a"
    u"\u7b2c\u4e00\uff0c\u4e00\u4e2a\u8dd1\u901a\u4e86\u7684\u4eff"
    u"\u771f\u63a7\u5236\u5668\uff0c\u5e76\u4e14\u5df2\u7ecf\u5728"
    u" GitHub \u4e0a\u5f00\u6e90\uff1b\u7b2c\u4e8c\uff0c"
    u"\u4e00\u4e2a\u521b\u65b0\u6027\u7684\u53ef\u7a7f\u6234\u786c"
    u"\u4ef6\u7ed3\u6784\u8bbe\u8ba1\uff1b\u7b2c\u4e09\uff0c"
    u"\u4e00\u6574\u5957\u5b8c\u6574\u7684\u6d88\u878d\u5b9e\u9a8c"
    u"\u8bbe\u8ba1\u548c\u5b9a\u91cf\u8bc4\u4f30\u6307\u6807\u3002",
    u"\u5728\u811a\u5e95\u677f\u7684\u8bbe\u8ba1\u4e0a\uff0c\u6211"
    u"\u4eec\u53c2\u8003\u4e86\u5b66\u672f\u754c\u51e0\u7bc7\u76f8"
    u"\u5173\u8bba\u6587\u7684\u65b9\u6848\uff0c\u6700\u7ec8\u786e"
    u"\u5b9a\u4f7f\u7528\u201c\u978b\u6258\u5f0f\u201d\u8bbe\u8ba1"
    u"\uff1a\u53d7\u8bd5\u8005\u76f4\u63a5\u7a7f\u7740\u81ea\u5df1"
    u"\u7684\u978b\u8e29\u8fdb\u978b\u6258\uff0c\u518d\u901a\u8fc7"
    u"\u7ed1\u5e26\u628a\u811a\u56fa\u5b9a\u4f4f\uff0c\u6574\u5957"
    u"\u7ed1\u5e26\u5f0f\u53ef\u7a7f\u6234\u88c5\u7f6e\u5e26\u52a8"
    u"\u8fd9\u4e2a\u978b\u6258\u8fd0\u52a8\u3002\u8fd9\u4e2a\u65b9"
    u"\u6848\u4e00\u65b9\u9762\u663e\u8457\u63d0\u5347\u4e86\u53d7"
    u"\u8bd5\u8005\u7684\u4f69\u6234\u820c\u9002\u5ea6\uff0c\u53e6"
    u"\u4e00\u65b9\u9762\u4e5f\u8ba9\u8bbe\u5907\u80fd\u591f\u901a"
    u"\u7528\u4e8e\u811a\u7801\u5dee\u5f02\u8f83\u5927\u7684\u4e0d"
    u"\u540c\u4eba\u7fa4\u3002",
    u"\u6d88\u878d\u5b9e\u9a8c\u4ee5\u53ca\u7528\u6765\u91cf\u5316"
    u"\u5224\u65ad\u8fd9\u4e2a\u88c5\u7f6e\u5230\u5e95\u662f\u5426"
    u"\u6709\u6548\u7684\u5eb7\u590d\u6307\u6807\u548c\u5b9e\u9a8c"
    u"\u6d41\u7a0b\uff0c\u6211\u4eec\u4e5f\u5df2\u7ecf\u5b8c\u6574"
    u"\u8bbe\u8ba1\u597d\u4e86\u3002",
    u"\u76ee\u524d\u8fd8\u6ca1\u6709\u9a8c\u8bc1\u4e24\u4ef6\u4e8b"
    u"\uff1a\u7b2c\u4e00\uff0c\u6574\u5957\u88c5\u7f6e\u5728\u72ec"
    u"\u7acb\u4f9b\u7535\u3001\u81ea\u7531\u53ef\u7a7f\u6234\u884c"
    u"\u8d70\u72b6\u6001\u4e0b\u7684\u8fd0\u884c\u8868\u73b0\uff1b"
    u"\u7b2c\u4e8c\uff0c\u771f\u5b9e\u53d7\u8bd5\u8005\u7684\u5b9e"
    u"\u9a8c\u6570\u636e\u3002\u524d\u671f\u5b9e\u9a8c\u5e73\u53f0"
    u"\u6211\u4eec\u5df2\u7ecf\u642d\u5efa\u597d\u5e76\u8dd1\u901a"
    u"\u4e86\uff0c\u4e0b\u4e00\u6b65\u5c31\u662f\u53d7\u8bd5\u8005"
    u"\u5b9e\u9a8c\u3002\u5982\u679c\u4e4b\u540e\u6709\u540c\u5b66"
    u"\u5bf9\u8fd9\u4e2a\u65b9\u5411\u611f\u5174\u8da3\uff0c\u53ef"
    u"\u4ee5\u5728\u6211\u4eec\u73b0\u5728\u8fd9\u4e2a\u4eff\u771f"
    u"\u4ee3\u7801\u5e93\u3001\u786c\u4ef6\u8bbe\u8ba1\u548c\u5b9e"
    u"\u9a8c\u534f\u8bae\u7684\u57fa\u7840\u4e0a\u76f4\u63a5\u5ef6"
    u"\u4f38\u4e0b\u53bb\u3002",
]

EN_BLOCKS = [
    "This is my final-term project. To set expectations honestly: the "
    "physical control loop has not yet been fully closed. What we have "
    "delivered are three concrete things \u2014 (i) a working simulation "
    "controller, which has been open-sourced on GitHub; (ii) an "
    "innovative wearable hardware design; and (iii) a complete "
    "experimental design with quantitative evaluation metrics.",
    "For the foot-plate, we surveyed several related ankle-rehab papers "
    "and converged on a \u201cshoe-tray\u201d design. The subject simply "
    "wears their own shoes and steps directly into the shoe tray, where "
    "straps fix the foot in place. Our wearable, strap-mounted device "
    "then drives the shoe tray. This scheme materially improves subject "
    "comfort, and \u2014 just as importantly \u2014 makes the device "
    "generalizable across users with very different foot sizes, which is "
    "a recurring pain point of fixed-geometry foot plates in the "
    "literature.",
    "The ablation experiments and the quantitative metrics that will be "
    "used to test whether the device produces clinically meaningful "
    "improvements have been fully designed; they are described in the "
    "previous section.",
    "Two things are still open. First, we have not yet validated the "
    "device's standalone operation \u2014 by which I mean battery-driven, "
    "free wearable walking, fully untethered. Second, we have not yet "
    "run the human-subject experiments themselves. The early experimental "
    "platform is built and runnable, and the open-source codebase, the "
    "wearable hardware design, and the experimental protocol together "
    "provide a turn-key starting point. If a future student is "
    "interested in this direction, they can extend this work directly "
    "from where we leave it today.",
]


def append(doc):
    doc.add_paragraph(ZH_HEADING, style="Heading 2")

    doc.add_paragraph(
        u"\U0001f1e8\U0001f1f3 \u4e2d\u6587\uff08\u53e3\u8bed\u7248\uff09",
        style="Heading 3",
    )
    for block in ZH_BLOCKS:
        doc.add_paragraph(block, style="normal")

    doc.add_paragraph(
        u"\U0001f1fa\U0001f1f8 English\uff08\u53e3\u8bed\u7248\uff09",
        style="Heading 3",
    )
    for block in EN_BLOCKS:
        doc.add_paragraph(block, style="normal")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--in-place",
        action="store_true",
        help="Also overwrite the original Roar.docx on disk.",
    )
    args = parser.parse_args()

    if not os.path.isfile(SOURCE):
        raise SystemExit("source not found: " + SOURCE)

    repo_dir = os.path.dirname(REPO_COPY)
    if not os.path.isdir(repo_dir):
        os.makedirs(repo_dir)

    doc = docx.Document(SOURCE)
    append(doc)
    doc.save(REPO_COPY)
    print("wrote " + REPO_COPY)

    if args.in_place:
        shutil.copyfile(REPO_COPY, SOURCE)
        print("also overwrote " + SOURCE)


if __name__ == "__main__":
    main()
